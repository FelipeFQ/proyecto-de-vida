import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = "/home/pipedev/Documents/Proyecto de vida/finanzas/clarita/Rendicion_Cuentas_La_Clarita.xlsx"

F = "Arial"
NAVY = "1B365D"
BLUE = "0000FF"
GREEN = "008000"
YELLOW = "FFFF00"
LIGHT = "EEF2F8"
GRAY = "D9D9D9"

CUR = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = 'dd/mm/yyyy'
MON = 'mmm yyyy'

thin = Side(style="thin", color=GRAY)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def q(sheet, cell):
    return f"'{sheet}'!{cell}"


wb = Workbook()
names = ["Resumen", "Propuesta", "Movimientos", "Cuadre", "Mobiliario", "Préstamo",
         "Registro Mensual", "Flujo Proyectado", "Escenarios", "Comparación", "Valoración", "Supuestos"]
ws0 = wb.active
ws0.title = names[0]
S = {names[0]: ws0}
for n in names[1:]:
    S[n] = wb.create_sheet(n)


def style(c, kind="f", bold=False, fmt=None, size=10, fill=None, wrap=False, align=None, italic=False):
    color = {"in": BLUE, "f": "000000", "link": GREEN, "hdr": "FFFFFF", "note": "555555"}[kind]
    c.font = Font(name=F, size=size, bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", start_color=fill)
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="center")


def put(ws, ref, val, kind="f", **kw):
    c = ws[ref]
    c.value = val
    style(c, kind, **kw)
    return c


def title(ws, text, sub=None, width="H"):
    put(ws, "A1", text, bold=True, size=15)
    ws["A1"].font = Font(name=F, size=15, bold=True, color=NAVY)
    if sub:
        put(ws, "A2", sub, "note", italic=True)
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, col=1, height=30):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=NAVY)
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = height


def section(ws, ref, text):
    c = put(ws, ref, text, bold=True, size=11)
    c.font = Font(name=F, size=11, bold=True, color=NAVY)


def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


def printsetup(ws, landscape=True, title_rows=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.center.text = "La Clarita — rendición de cuentas · Página &P de &N"
    ws.oddFooter.center.size = 8


def confirm(c, text):
    c.fill = PatternFill("solid", start_color=YELLOW)
    c.comment = Comment(text, "Felipe")


# =====================================================================
# SUPUESTOS
# =====================================================================
ws = S["Supuestos"]
SP = "Supuestos"
title(ws, "Supuestos y parámetros",
      "Azul = dato ingresado · Negro = fórmula · Verde = viene de otra hoja · Amarillo = POR CONFIRMAR. Cambie aquí y todo el libro se recalcula.")
section(ws, "A4", "Habitaciones")
header(ws, 5, ["Habitación", "Tipo", "Canon mensual", "Fecha ingreso inquilino", "Estado",
               "Fecha compra mobiliario", "Mes 1ª cuota muebles", "Mes última cuota", "Costo amoblado"], height=32)
rooms = [
    ("Hab 1", "Grande", 1250000, dt.date(2026, 9, 15), "Contrato firmado — inicio 15 sep 2026", dt.date(2026, 9, 1)),
    ("Hab 2", "Grande", 1500000, dt.date(2026, 10, 15), "Interesado aceptó precio — reuniendo documentos", dt.date(2026, 10, 15)),
    ("Hab 3", "Grande", 1350000, dt.date(2026, 11, 1), "Proyectado (precio mínimo esperado)", dt.date(2026, 11, 1)),
    ("Hab 4", "Grande", 1350000, dt.date(2026, 12, 1), "Proyectado (precio mínimo esperado)", dt.date(2026, 12, 1)),
    ("Hab 5", "Pequeña", 1100000, dt.date(2027, 1, 1), "Proyectado", dt.date(2027, 1, 1)),
    ("Hab 6", "Pequeña", 1100000, dt.date(2027, 2, 1), "Proyectado", dt.date(2027, 2, 1)),
    ("Hab 7", "Grande", 1400000, dt.date(2028, 4, 1), "Hoy es bodega de objetos de los abuelos — disponible en 1–2 años", dt.date(2028, 4, 1)),
]
mob_cols = ["B", "C", "D", "E", "F", "G", "H"]
for i, (h, t, canon, fi, est, fc) in enumerate(rooms):
    r = 6 + i
    put(ws, f"A{r}", h, bold=True)
    put(ws, f"B{r}", t, "in")
    put(ws, f"C{r}", canon, "in", fmt=CUR)
    put(ws, f"D{r}", fi, "in", fmt=DATE)
    put(ws, f"E{r}", est, "in")
    put(ws, f"F{r}", fc, "in", fmt=DATE)
    put(ws, f"G{r}", f"=DATE(YEAR(F{r}),MONTH(F{r}),1)", fmt=MON)
    put(ws, f"H{r}", f"=EDATE(G{r},$B$43-1)", fmt=MON)
    put(ws, f"I{r}", f"={q('Mobiliario', mob_cols[i] + '44')}", "link", fmt=CUR)
confirm(ws["D12"], "Supuesto base: abril 2028 (punto medio de 'en 1 o 2 años'). Cambie la fecha si hay más claridad.")
confirm(ws["C12"], "Felipe: entre 1.300.000 y 1.500.000. Se usa el punto medio.")
put(ws, "A13", "Total 7 habitaciones", bold=True)
put(ws, "C13", "=SUM(C6:C12)", bold=True, fmt=CUR)
put(ws, "I13", "=SUM(I6:I12)", bold=True, fmt=CUR)
put(ws, "A14", "Habitaciones 3 a 7: precios esperados por Felipe. Ritmo base: 1 habitación nueva por mes hasta la 6. La 7 se habilita cuando salgan los objetos almacenados (se trasladan o se venden).", "note", italic=True)

section(ws, "A15", "Gastos mensuales de operación")
header(ws, 16, ["Concepto", "Actual (1 inquilino)", "Estimado (6 inquilinos)", "Nota"])
gastos = [
    ("Energía", 170000, 600000, "7 duchas eléctricas a 220V: el rubro más sensible"),
    ("Gas", 46000, 120000, ""),
    ("Acueducto", 80000, 220000, ""),
    ("Internet", 200000, 200000, ""),
    ("Aseo", 110000, 500000, "Persona de aseo semanal — valor mensual acordado"),
]
for i, (c_, a, e, n) in enumerate(gastos):
    r = 17 + i
    put(ws, f"A{r}", c_)
    put(ws, f"B{r}", a, "in", fmt=CUR)
    put(ws, f"C{r}", e, "in", fmt=CUR)
    put(ws, f"D{r}", n, "note")
put(ws, "A22", "Subtotal servicios", bold=True)
put(ws, "B22", "=SUM(B17:B21)", bold=True, fmt=CUR)
put(ws, "C22", "=SUM(C17:C21)", bold=True, fmt=CUR)
put(ws, "A23", "Reserva reparaciones menores/mayores + vacancia")
put(ws, "B23", 0, "in", fmt=CUR)
put(ws, "C23", 1300000, "in", fmt=CUR)
put(ws, "D23", "Se provisiona en proporción a las habitaciones ocupadas", "note")
put(ws, "A24", "Ahorro mensual impuesto predial")
put(ws, "B24", 100000, "in", fmt=CUR)
put(ws, "C24", 100000, "in", fmt=CUR)
put(ws, "A25", "Total gastos de operación", bold=True)
put(ws, "B25", "=B22+B23+B24", bold=True, fmt=CUR)
put(ws, "C25", "=C22+C23+C24", bold=True, fmt=CUR)
put(ws, "D25", "Fuente: estimados de Felipe (sep 2026). Sin la comisión de administración.", "note")

section(ws, "A27", "Parámetros financieros")
params = [
    # row, label, value, fmt, kind, note
    (28, "Comisión de administración seleccionada (sobre ingreso neto operativo)", 0.20, PCT, "in", "Cambie a 10%, 20% o 30% para ver cada escenario en todo el libro"),
    (29, "Escenario comisión 1", 0.10, PCT, "in", ""),
    (30, "Escenario comisión 2", 0.20, PCT, "in", ""),
    (31, "Escenario comisión 3", 0.30, PCT, "in", ""),
    (32, "Préstamo mamá — monto", 35000000, CUR, "in", "Fuente: Felipe"),
    (33, "Préstamo mamá — tasa nominal mensual", 0.015, '0.00%', "in", "Fuente: Felipe"),
    (34, "Préstamo mamá — fecha de desembolso (para cálculo de intereses)", dt.date(2026, 8, 15), DATE, "in", "Acordado: intereses desde el 15 ago, aunque hubo desembolsos parciales desde el 15 jul"),
    (35, "Préstamo mamá — fecha primera cuota", dt.date(2026, 12, 15), DATE, "in", "Fuente: Felipe"),
    (36, "Préstamo mamá — número de cuotas (sistema francés)", 36, '0', "in", "3 años desde la primera cuota"),
    (37, "Préstamo mamá — tasa efectiva anual equivalente", "=(1+B33)^12-1", PCT, "f", ""),
    (38, "Cuenta de ahorros abuelos — rendimiento efectivo anual", 0.085, PCT, "in", "Fuente: Felipe"),
    (39, "Cuenta de ahorros abuelos — saldo (ILUSTRATIVO)", 100000000, CUR, "in", "Valor ilustrativo, no es el saldo real"),
    (40, "CDT 360 días — tasa promedio del mercado (EA)", 0.1207, PCT, "in", "Banco de la República vía Infobae/El País, sep 2026"),
    (41, "CDT 360 días — tasa alta del mercado (EA)", 0.13, PCT, "in", "KOA / Pibank, sep 2026 (finanzasplus.co)"),
    (42, "Rendimiento mensual equivalente de la cuenta de ahorros", "=(1+B38)^(1/12)-1", '0.000%', "f", ""),
    (43, "Mobiliario — número de cuotas por habitación (derivado de la cuota promedio, celda F34)", "=ROUND(Mobiliario!B44/F34,0)", '0', "f", "Cuotas sin interés de Mercado Libre"),
    (44, "Mobiliario — costo habitación pequeña vs. grande", 1.0, PCT, "in", "Felipe: rango parecido"),
    (45, "Lavandería nueva — lavadora + monedero", "=Mobiliario!C57", CUR, "link", "Se elige en la hoja Mobiliario"),
    (46, "Lavadora nueva — mes de compra", dt.date(2026, 10, 1), MON, "in", ""),
    (47, "Lavadora nueva — ahorro mensual estimado (energía + agua)", 45000, CUR, "in", "Estimado de Felipe: entre 40.000 y 50.000"),
    (50, "Ingreso neto de la casa ANTES de la remodelación (mensual)", 620000, CUR, "in", "Felipe: bruto ~820.000, neto ~620.000"),
    (52, "Referencia de mercado: honorario de administración inmobiliaria (sobre canon bruto)", 0.10, PCT, "in", "Bogotá 8%–12% + IVA (Inmobiliare Latam, guía 2026)"),
    (53, "IVA sobre honorarios", 0.19, PCT, "in", ""),
    (54, "Primer mes de la proyección", dt.date(2026, 10, 1), MON, "in", ""),
]
for r, lab, val, fmt, kind, note in params:
    put(ws, f"A{r}", lab, wrap=True)
    put(ws, f"B{r}", val, kind, fmt=fmt)
    put(ws, f"C{r}", note, "note")
for r in (39,):
    confirm(ws[f"B{r}"], ws[f"C{r}"].value)
section(ws, "E27", "Habitación 7, lavandería y muebles")
newp = [
    (28, "Hab 7 — venta de objetos almacenados (ingreso único)", 0, CUR, "in", "Valor desconocido, probablemente < $10M. Se deja en 0 (conservador)"),
    (29, "Monedero — precio por ciclo de lavado", 7000, CUR, "in", "Definido por Felipe. El contrato ya contempla cobro por ciclo (gratis hasta instalar el monedero)"),
    (30, "Ciclos de lavado por inquilino al mes", 4, '0', "in", "Supuesto: 1 por semana"),
    (31, "Monedero — mes de inicio", "=B46", MON, "f", "Se compra a cuotas sin interés junto con la lavadora"),
    (32, "¿Se incluye el monedero? (1 = sí, 0 = no)", "=Mobiliario!B56", '0', "link", "Se elige en la hoja Mobiliario"),
    (33, "Inquilinos del estimado de gastos (columna C)", 6, '0', "in", "Los gastos se extrapolan por inquilino adicional"),
    (34, "Cuota mensual promedio de muebles por habitación", 850000, CUR, "in", "Felipe: desde el 2º mes ~850.000 por habitación"),
    (35, "Mes en que entra la venta de objetos (= compra muebles Hab 7)", "=G12", MON, "f", ""),
]
for r, lab, val, fmt, kind, note in newp:
    put(ws, f"E{r}", lab, wrap=True)
    put(ws, f"F{r}", val, kind, fmt=fmt)
    put(ws, f"G{r}", note, "note")

section(ws, "A56", "Listas (usadas en Movimientos)")
header(ws, 57, ["Categoría", "Grupo"])
cats = [
    ("Obra – contratista principal", "Obra y adecuaciones"),
    ("Costos bancarios – 4x1000", "Obra y adecuaciones"),
    ("Eléctrico – aumento de carga (Enel)", "Obra y adecuaciones"),
    ("Seguridad – cámaras", "Obra y adecuaciones"),
    ("Hidráulico – bomba de presión", "Obra y adecuaciones"),
    ("Mobiliario – habitaciones", "Mobiliario"),
    ("Lavandería – lavadora", "Mobiliario"),
    ("Traslado entre cuentas", "Traslado (no es gasto)"),
]
for i, (c_, g) in enumerate(cats):
    put(ws, f"A{58 + i}", c_, "in")
    put(ws, f"B{58 + i}", g, "in")
header(ws, 57, ["Fuente del dinero"], col=4)
fuentes = ["Abuelos – capital propio", "Préstamo mamá", "TC abuelos (cuotas)", "Arriendo de la casa",
           "Felipe (por definir)", "Por confirmar"]
for i, f_ in enumerate(fuentes):
    put(ws, f"D{58 + i}", f_, "in")
header(ws, 57, ["Estado del soporte"], col=6)
estados = ["Soportado", "Por desglosar", "Pendiente de pago"]
for i, e in enumerate(estados):
    put(ws, f"F{58 + i}", e, "in")
widths(ws, {"A": 58, "B": 18, "C": 16, "D": 22, "E": 40, "F": 20, "G": 14, "H": 14, "I": 16})
printsetup(ws)

# =====================================================================
# MOVIMIENTOS (ledger)
# =====================================================================
ws = S["Movimientos"]
title(ws, "Movimientos — cada pago, uno por uno",
      "Documento vivo: agregue una fila por cada pago, recibo firmado o transferencia. Las filas 'Por desglosar' se reemplazan por sus pagos individuales cuando se digitalicen los soportes.")
put(ws, "A3", "Cómo llenar: escriba en azul · Categoría, Fuente y Estado tienen lista desplegable · Grupo se calcula solo · Tipo: 'Gasto' o 'Traslado'.", "note")
hdr = ["N°", "Fecha", "Categoría", "Concepto / detalle", "Proveedor o beneficiario", "Fuente del dinero",
       "Quién pagó (cuenta / persona)", "Medio", "Monto", "Soporte (tipo y referencia)", "Estado", "Tipo", "Grupo", "Notas"]
header(ws, 5, hdr, height=34)
L0, L1 = 6, 205  # ledger rows
D = dt.date
GS = "Germán Segura Rojas (CC 80.127.526)"
GC = "Clemente Gregorio Carriazo"
OH = "Oscar Mauricio Herrera Bello"
AB = "Abuelos – capital propio"
PM = "Préstamo mamá"
OB = "Obra – contratista principal"
EL = "Eléctrico – aumento de carga (Enel)"
CAM = "Seguridad – cámaras"
EFE = "Efectivo retirado por el abuelo, entregado por Felipe"
mov = [
    (D(2026, 5, 15), OB, "Anticipo de obra (transferencia 1 de 2)", GS, AB, "Jorge Enrique Quintero Ruiz — Nu ••8689", "Transferencia", 10000000, "Comprobante Nu 22848468…923405", "Soportado", "Gasto", ""),
    (D(2026, 5, 15), OB, "Anticipo de obra (transferencia 2 de 2)", GS, AB, "Jorge Enrique Quintero Ruiz — Nu ••8689", "Transferencia", 10000000, "Comprobante Nu 10341941…199155", "Soportado", "Gasto", ""),
    (D(2026, 5, 15), "Costos bancarios – 4x1000", "Impuesto 4x1000 de las dos transferencias del anticipo", "DIAN (vía Nu)", AB, "Jorge Enrique Quintero Ruiz — Nu ••8689", "Débito automático", 80000, "Mismos comprobantes Nu", "Soportado", "Gasto", "$40.000 por cada transferencia"),
    (D(2026, 5, 25), OB, "Adelanto obra La Clarita", GS, AB, EFE, "Efectivo", 14000000, "Recibo firmado 25/05/2026", "Soportado", "Gasto", ""),
    (D(2026, 6, 22), OB, "Avance de obra La Clarita", GS, AB, EFE, "Efectivo", 15000000, "Recibo firmado 22/06/2026", "Soportado", "Gasto", ""),
    (D(2026, 6, 22), EL, "50% adecuación eléctrica", GC, AB, EFE, "Efectivo", 2500000, "Recibo firmado 22/06/2026", "Soportado", "Gasto", ""),
    (D(2026, 6, 24), OB, "Avance obra La Clarita", GS, AB, EFE, "Efectivo", 10500000, "Recibo firmado 24/06/2026", "Soportado", "Gasto", ""),
    (D(2026, 7, 15), OB, "Pago de obra (adicionales)", GS, PM, "Janneth Quintero Ávila — Nu", "Transferencia", 10000000, "Comprobante Nu 21299625…062555", "Soportado", "Gasto", ""),
    (D(2026, 7, 15), OB, "Pago de obra (adicionales)", GS, PM, "Janneth Quintero Ávila — Nu", "Transferencia", 8000000, "Comprobante Nu 28378064…566602", "Soportado", "Gasto", ""),
    (D(2026, 7, 25), EL, "Segundo pago adecuación eléctrica", GC, AB, "Andrés Felipe Fajardo — Nu ••6045 (dinero del abuelo)", "Transferencia", 1250000, "Comprobante Nu 13857091…720148", "Soportado", "Gasto", ""),
    (D(2026, 8, 13), OB, "Pago de obra (adicionales)", GS, PM, "Janneth Quintero Ávila — Nu", "Transferencia", 10500000, "Comprobante Nu 58860168…806189", "Soportado", "Gasto", ""),
    (D(2026, 7, 15), "Costos bancarios – 4x1000", "Impuesto 4x1000 de las transferencias del 15 jul y 13 ago", "DIAN (vía Nu)", PM, "Janneth Quintero Ávila — Nu", "Débito automático", 16188, "Mismos comprobantes Nu", "Soportado", "Gasto", "$4.211 + $11.977. Asumido por la cuenta de mamá, adicional a los $35M"),
    (D(2026, 9, 7), OB, "Abono de obra", GS, AB, "Jorge Enrique Quintero — Davivienda ahorros ••5040", "Transferencia", 15000000, "Davivienda, aprobación 10420591", "Soportado", "Gasto", "Destino: cuenta AV Villas ••3613 de Germán Segura"),
    (D(2026, 9, 7), OB, "Abono 1er contrato", GS, AB, EFE, "Efectivo", 5000000, "Recibo firmado 07/09/2026", "Soportado", "Gasto", ""),
    (D(2026, 9, 9), CAM, "Cámaras — pago 1", OH, PM, "Janneth Quintero Ávila — Nu ••0510", "Transferencia", 2221135, "Comprobante Nu 10234145…242438", "Soportado", "Gasto", ""),
    (D(2026, 9, 15), OB, "Abono 1er contrato", GS, AB, EFE, "Efectivo", 5000000, "Recibo firmado 15/09/2026", "Soportado", "Gasto", ""),
    (D(2026, 9, 29), OB, "Pago final de obra", GS, PM, "Andrés Felipe Fajardo — Nu ••6045 (dinero del préstamo transferido por mamá)", "Transferencia", 1000000, "Comprobante Nu 10694692…366546", "Soportado", "Gasto", ""),
    (D(2026, 9, 29), CAM, "Cámaras — pago 2", OH, PM, "Andrés Felipe Fajardo — Nu ••6045 (dinero del préstamo transferido por mamá)", "Transferencia", 1778650, "Comprobante Nu 10893610…120247", "Soportado", "Gasto", ""),
    (D(2026, 9, 30), CAM, "Cámaras — pago 3", OH, PM, "Janneth Quintero Ávila — Nu ••0510", "Transferencia", 1000000, "Comprobante Nu 39673969…293348", "Soportado", "Gasto", "Total cámaras = cotización final $4.999.785"),
    (None, "Traslado entre cuentas", "Sobrante del préstamo trasladado a la cuenta de la abuela (cuenta de la casa)", "Cuenta abuela", PM, "Cuenta de mamá", "Transferencia", 500215, "Consultable en la app de la cuenta de la abuela", "Soportado", "Traslado", "Saldo = $35.000.000 − $29.500.000 − $4.999.785"),
    (None, "Hidráulico – bomba de presión", "Bomba de presurización (equipo)", "Por confirmar", "TC abuelos (cuotas)", "Tarjeta de crédito abuelo", "Tarjeta de crédito", 310000, "Consultable en la app de la tarjeta del abuelo", "Soportado", "Gasto", ""),
    (D(2026, 10, 1), "Hidráulico – bomba de presión", "Instalación bomba (mano de obra)", "Instalador", AB, "Abuelo", "Por definir", 300000, "Pendiente", "Pendiente de pago", "Gasto", "Instalación programada 1 oct 2026"),
    (None, EL, "Saldo adecuación eléctrica", GC, AB, "Abuelo", "Por definir", 1250000, "Pendiente", "Pendiente de pago", "Gasto", "Obra civil terminada. Se paga cuando Enel cambie el contador y apruebe el aumento de carga"),
    (None, "Mobiliario – habitaciones", "Amoblado completo habitación 1 (37 ítems — ver hoja Mobiliario)", "Mercado Libre (varias tiendas)", "TC abuelos (cuotas)", "Tarjeta de crédito abuelos", "Tarjeta de crédito", "='Mobiliario'!B42", "Compras consultables en Mercado Libre y en la app de la tarjeta", "Soportado", "Gasto", "Cuotas sin interés 6–12 meses. 1ª cuota (~$1.300.000) pagada con el arriendo de sep 2026"),
]
for i, row in enumerate(mov):
    r = L0 + i
    put(ws, f"A{r}", i + 1)
    vals = list(row)
    for j, v in enumerate(vals):
        col = "BCDEFGHIJKLN"[j]
        c = put(ws, f"{col}{r}", v, "in", wrap=(col in "DGJN"))
        if col == "B":
            c.number_format = DATE
        if col == "I":
            c.number_format = CUR
for r in range(L0, L1 + 1):
    put(ws, f"M{r}", f'=IF(C{r}="","",IFERROR(VLOOKUP(C{r},{q(SP, "$A$58:$B$65")},2,FALSE),"Revisar"))')
    if r >= L0 + len(mov):
        ws[f"A{r}"].value = f'=IF(C{r}="","",ROW()-{L0 - 1})'
        style(ws[f"A{r}"])
        for col in "BCDEFGHIJKLN":
            style(ws[f"{col}{r}"], "in", fmt=(DATE if col == "B" else CUR if col == "I" else None))
    for col in "ABCDEFGHIJKLMN":
        ws[f"{col}{r}"].border = BOX
for i, row in enumerate(mov):
    if "POR CONFIRMAR" in row[5] or "confirmar" in row[11]:
        confirm(ws[f"G{L0 + i}"], row[11] or "Confirmar cuenta de origen")
    if row[4] == "Por confirmar":
        confirm(ws[f"F{L0 + i}"], "¿Quién paga?")

dv_cat = DataValidation(type="list", formula1=f"={q(SP, '$A$58:$A$65')}", allow_blank=True)
dv_src = DataValidation(type="list", formula1=f"={q(SP, '$D$58:$D$63')}", allow_blank=True)
dv_est = DataValidation(type="list", formula1=f"={q(SP, '$F$58:$F$60')}", allow_blank=True)
dv_tipo = DataValidation(type="list", formula1='"Gasto,Traslado"', allow_blank=True)
for dv, col in ((dv_cat, "C"), (dv_src, "F"), (dv_est, "K"), (dv_tipo, "L")):
    ws.add_data_validation(dv)
    dv.add(f"{col}{L0}:{col}{L1}")
put(ws, "H4", "Total gastos registrados:", bold=True, align="right")
put(ws, "I4", f'=SUMIFS(I{L0}:I{L1},L{L0}:L{L1},"Gasto")', bold=True, fmt=CUR)
widths(ws, {"A": 5, "B": 11, "C": 26, "D": 36, "E": 22, "F": 20, "G": 26, "H": 15, "I": 14, "J": 28, "K": 14, "L": 9, "M": 18, "N": 38})
ws.freeze_panes = "C6"
printsetup(ws, title_rows="5:5")
ws.print_area = f"A1:N{L0 + 40}"

MV = "Movimientos"
M_I = q(MV, f"$I${L0}:$I${L1}")
M_C = q(MV, f"$C${L0}:$C${L1}")
M_F = q(MV, f"$F${L0}:$F${L1}")
M_K = q(MV, f"$K${L0}:$K${L1}")
M_L = q(MV, f"$L${L0}:$L${L1}")
M_M = q(MV, f"$M${L0}:$M${L1}")

# =====================================================================
# MOBILIARIO
# =====================================================================
ws = S["Mobiliario"]
title(ws, "Mobiliario por habitación",
      "Hab 1 = costo real pagado. Hab 2–7: escriba el costo real de cada ítem al comprarlo. Mientras una habitación no tenga datos, se usa el costo de la Hab 1 como estimado.")
header(ws, 4, ["Ítem", "Hab 1 (real)", "Hab 2", "Hab 3", "Hab 4", "Hab 5", "Hab 6", "Hab 7"])
items = [("Ducha", 172200), ("Gabinete de baño", 139900), ("Bowl acero inoxidable", 12400), ("Tijeras de cocina", 15600),
         ("Estantería de ducha", 28000), ("Basurero de pedal 20 litros", 70500), ("Microondas Electrolux", 250000),
         ("Vajilla 2 puestos Corona", 52000), ("Set de cubiertos en acero", 20000), ("Nevera minibar", 620000),
         ("Filtro de agua ozono", 180000), ("Cobijas x2", 80000), ("Cesta de ropa sucia", 64000),
         ("Set de utensilios de cocina", 54500), ("Batería de ollas Imusa", 210000), ("Set de cuchillos Tramontina", 90000),
         ("Tope de puertas x2", 8000), ("Silla de oficina", 355000), ("Papelera de baño 3 litros", 55000),
         ("Licuadora Universal", 93000), ("Vasos de cristal x2", 11000), ("Juego de sábanas x2", 82000),
         ("Mesa de noche", 145000), ("Televisor 32\" Challenger", 492000), ("Escritorio 120 cm", 204000),
         ("Colchón pullman 120 cm", 1300000), ("Basecama semidoble", 340000), ("Clóset 120 cm", 740000),
         ("Almohadas x2", 138000), ("Organizador de cubiertos", 74000), ("Tabla para picar de acero", 24500),
         ("Soporte de TV", 42000), ("Protector de colchón", 60000), ("Lámpara de mesa de noche", 60200),
         ("Lámpara de escritorio", 60000), ("Timbre", 40000)]
for i, (it, v) in enumerate(items):
    r = 5 + i
    put(ws, f"A{r}", it)
    put(ws, f"B{r}", v, "in", fmt=CUR)
    for col in "CDEFGH":
        style(ws[f"{col}{r}"], "in", fmt=CUR)
    for col in "ABCDEFGH":
        ws[f"{col}{r}"].border = BOX
last = 5 + len(items) - 1  # 40
put(ws, "A41", "Mesa abatible multiusos para cocina")
ws["B41"].value = 169000
for col in "BCDEFGH":
    style(ws[f"{col}41"], "in", fmt=CUR)
put(ws, "A42", "Total registrado (real)", bold=True)
put(ws, "A43", "¿Real o estimado?", bold=True)
put(ws, "A44", "Costo usado en los cálculos", bold=True)
for k, col in enumerate("BCDEFGH"):
    put(ws, f"{col}42", f"=SUM({col}5:{col}41)", bold=True, fmt=CUR)
    put(ws, f"{col}43", f'=IF({col}42>0,"Real","Estimado")', align="center")
    if col == "B":
        put(ws, "B44", "=B42", bold=True, fmt=CUR)
    else:
        tipo_row = 6 + k
        put(ws, f"{col}44",
            f'=IF({col}42>0,{col}42,$B$42*IF({q(SP, f"$B${tipo_row}")}="Pequeña",{q(SP, "$B$44")},1))', bold=True, fmt=CUR)
    ws[f"{col}44"].fill = PatternFill("solid", start_color=LIGHT)
put(ws, "A45", "Total mobiliario 7 habitaciones", bold=True)
put(ws, "B45", "=SUM(B44:H44)", bold=True, fmt=CUR)


section(ws, "A48", "Lavandería — lavadora nueva y monedero")
header(ws, 49, ["Opción", "Precio", "Nota"])
put(ws, "A50", "Whirlpool 14 kg")
put(ws, "B50", 2830000, "in", fmt=CUR)
put(ws, "C50", "Cotizada por Felipe", "note")
put(ws, "A51", "Whirlpool 20 kg")
put(ws, "B51", 3350000, "in", fmt=CUR)
put(ws, "C51", "Cotizada por Felipe — permite lavar cobijas y edredones", "note")
put(ws, "A52", "Opción seleccionada (escriba 1 = 14 kg, 2 = 20 kg)", bold=True)
put(ws, "B52", 1, "in")
put(ws, "C52", "=INDEX(B50:B51,B52)", bold=True, fmt=CUR)
put(ws, "A54", "Monedero para cobro por ciclo")
put(ws, "B54", 1190000, "in", fmt=CUR)
put(ws, "C54", "Cotizado por Felipe — ordena el uso y genera ingreso por ciclo", "note")
put(ws, "A56", "¿Incluir monedero? (1 = sí, 0 = no)", bold=True)
put(ws, "B56", 1, "in")
put(ws, "A57", "TOTAL LAVANDERÍA", bold=True)
put(ws, "C57", "=C52+B54*B56", bold=True, fmt=CUR)
put(ws, "A59", "Lavadora actual: Whirlpool de más de 20 años, alto consumo de energía. Zonas comunes ya amobladas (a mediano plazo: puffs o sillas de lectura).", "note", italic=True)
widths(ws, {"A": 44, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15})
ws.freeze_panes = "B5"
printsetup(ws, landscape=False, title_rows="4:4")

# =====================================================================
# PRÉSTAMO
# =====================================================================
ws = S["Préstamo"]
title(ws, "Préstamo de mamá — tabla de amortización (sistema francés)",
      "Intereses de gracia se suman al capital; cuota fija desde el 15 dic 2026. Columnas J–K: registre cada pago real.")
lp = [
    (3, "Monto prestado", f"={q(SP, '$B$32')}", CUR, "link"),
    (4, "Tasa nominal mensual", f"={q(SP, '$B$33')}", '0.00%', "link"),
    (5, "Tasa efectiva anual equivalente", "=(1+B4)^12-1", PCT, "f"),
    (6, "Fecha de desembolso", f"={q(SP, '$B$34')}", DATE, "link"),
    (7, "Fecha primera cuota", f"={q(SP, '$B$35')}", DATE, "link"),
    (8, "Número de cuotas", f"={q(SP, '$B$36')}", '0', "link"),
    (9, "Meses de interés sumado al capital antes del primer periodo", "=(YEAR(B7)-YEAR(B6))*12+MONTH(B7)-MONTH(B6)-1", '0', "f"),
    (10, "Saldo sobre el que se calculan las cuotas", "=B3*(1+B4)^B9", CUR, "f"),
    (11, "Cuota fija mensual", "=PMT(B4,B8,-B10)", CUR, "f"),
    (12, "Total a pagar en 36 cuotas", "=B11*B8", CUR, "f"),
    (13, "Total intereses si se paga en cuotas", "=B12-B3", CUR, "f"),
]
for r, lab, f_, fmt, kind in lp:
    put(ws, f"A{r}", lab)
    put(ws, f"B{r}", f_, kind, fmt=fmt, bold=(r in (11, 13)))
header(ws, 17, ["N°", "Fecha", "Mes", "Tipo", "Saldo inicial", "Interés", "Cuota", "Abono a capital", "Saldo final",
                "Pagado real", "Fecha pago real", "Nota"])
for k in range(40):
    r = 18 + k
    put(ws, f"A{r}", k)
    put(ws, f"B{r}", f"=EDATE($B$6,A{r})", fmt=DATE)
    put(ws, f"C{r}", f"=DATE(YEAR(B{r}),MONTH(B{r}),1)", fmt=MON)
    put(ws, f"D{r}", f'=IF(A{r}=0,"Desembolso",IF(B{r}<$B$7,"Gracia (interés se suma)","Cuota "&(A{r}-$B$9)))')
    put(ws, f"E{r}", 0 if k == 0 else f"=I{r - 1}", fmt=CUR)
    put(ws, f"F{r}", 0 if k == 0 else f"=E{r}*$B$4", fmt=CUR)
    put(ws, f"G{r}", 0 if k == 0 else f"=IF(B{r}>=$B$7,$B$11,0)", fmt=CUR)
    put(ws, f"H{r}", f"=G{r}-F{r}", fmt=CUR)
    put(ws, f"I{r}", "=$B$3" if k == 0 else f"=MAX(0,ROUND(E{r}+F{r}-G{r},0))", fmt=CUR)
    style(ws[f"J{r}"], "in", fmt=CUR)
    style(ws[f"K{r}"], "in", fmt=DATE)
    for col in "ABCDEFGHIJKL":
        ws[f"{col}{r}"].border = BOX
widths(ws, {"A": 52, "B": 16, "C": 11, "D": 22, "E": 15, "F": 13, "G": 13, "H": 15, "I": 15, "J": 13, "K": 13, "L": 24})
ws.column_dimensions["A"].width = 52
printsetup(ws, title_rows="17:17")
PR = "Préstamo"

# =====================================================================
# FLUJO PROYECTADO
# =====================================================================
ws = S["Flujo Proyectado"]
FP = "Flujo Proyectado"
title(ws, "Flujo proyectado mes a mes — 72 meses",
      "Llenado gradual según fechas de ingreso (hoja Supuestos). En ambas opciones el arriendo paga la operación y la cuota del préstamo de mamá. La diferencia es de dónde salen las cuotas de los muebles.")
groups = [("A4:I4", "OPERACIÓN DE LA CASA"), ("J4:M4", "OPCIÓN A — el arriendo paga también los muebles"),
          ("N4:Q4", "OPCIÓN B — los ahorros pagan los muebles"), ("R4:T4", "RECUPERACIÓN ACUMULADA POR COMISIÓN")]
for rng, t in groups:
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = t
    c.font = Font(name=F, size=10, bold=True, color=NAVY)
    c.fill = PatternFill("solid", start_color=LIGHT)
    c.alignment = Alignment(horizontal="center")
fh = ["Mes", "Hab. ocupadas", "Ingreso bruto (arriendos + lavandería)", "Servicios", "Reserva reparac. + vacancia", "Ahorro predial",
      "Neto operativo", "Comisión administración", "Neto casa (antes de deudas)",
      "Cuota préstamo mamá", "Cuotas muebles + lavadora", "Caja casa A", "Acumulado A",
      "Caja casa B", "Acumulado B", "Tomado de ahorros (muebles − venta objetos)", "Acumulado tomado de ahorros",
      "Acum. neto 10%", "Acum. neto 20%", "Acum. neto 30%"]
header(ws, 5, fh, height=44)
R0, R1 = 6, 77
sp = lambda c: q(SP, c)
for k in range(72):
    r = R0 + k
    put(ws, f"A{r}", f"={sp('$B$54')}" if k == 0 else f"=EDATE(A{r - 1},1)", "link" if k == 0 else "f", fmt=MON)
    put(ws, f"B{r}", f"=SUMPRODUCT(--({sp('$D$6:$D$12')}<=A{r}))", fmt='0')
    put(ws, f"C{r}", f"=SUMPRODUCT(({sp('$D$6:$D$12')}<=A{r})*{sp('$C$6:$C$12')})+IF(A{r}>={sp('$F$31')},B{r}*{sp('$F$30')}*{sp('$F$29')}*{sp('$F$32')},0)", fmt=CUR)
    put(ws, f"D{r}", f"=IF(B{r}=0,0,{sp('$B$22')}+({sp('$C$22')}-{sp('$B$22')})*(B{r}-1)/({sp('$F$33')}-1))-IF(A{r}>={sp('$B$46')},{sp('$B$47')},0)", fmt=CUR)
    put(ws, f"E{r}", f"={sp('$C$23')}*B{r}/{sp('$F$33')}", fmt=CUR)
    put(ws, f"F{r}", f"={sp('$C$24')}", fmt=CUR)
    put(ws, f"G{r}", f"=C{r}-D{r}-E{r}-F{r}", fmt=CUR)
    put(ws, f"H{r}", f"=MAX(0,G{r})*{sp('$B$28')}", fmt=CUR)
    put(ws, f"I{r}", f"=G{r}-H{r}", fmt=CUR, bold=True)
    put(ws, f"J{r}", f"=IFERROR(INDEX({q(PR, '$G$18:$G$57')},MATCH(A{r},{q(PR, '$C$18:$C$57')},0)),0)", fmt=CUR)
    lav_end = f"EDATE({sp('$B$46')},{sp('$B$43')}-1)"
    put(ws, f"K{r}",
        f"=SUMPRODUCT(({sp('$G$6:$G$12')}<=A{r})*({sp('$H$6:$H$12')}>=A{r})*{sp('$I$6:$I$12')})/{sp('$B$43')}"
        f"+IF(AND(A{r}>={sp('$B$46')},A{r}<={lav_end}),{sp('$B$45')}/{sp('$B$43')},0)", fmt=CUR)
    put(ws, f"L{r}", f"=I{r}-J{r}-K{r}", fmt=CUR, bold=True)
    put(ws, f"M{r}", f"=L{r}" if k == 0 else f"=M{r - 1}+L{r}", fmt=CUR)
    put(ws, f"N{r}", f"=I{r}-J{r}", fmt=CUR, bold=True)
    put(ws, f"O{r}", f"=N{r}" if k == 0 else f"=O{r - 1}+N{r}", fmt=CUR)
    put(ws, f"P{r}", f"=K{r}-IF(A{r}={sp('$F$35')},{sp('$F$28')},0)", fmt=CUR)
    put(ws, f"Q{r}", f"=P{r}" if k == 0 else f"=Q{r - 1}+P{r}", fmt=CUR)
    for col, prow in (("R", "$B$29"), ("S", "$B$30"), ("T", "$B$31")):
        base = f"(G{r}-MAX(0,G{r})*{sp(prow)})"
        put(ws, f"{col}{r}", f"={base}" if k == 0 else f"={col}{r - 1}+{base}", fmt=CUR)
    if k % 12 == 11:
        for col in "ABCDEFGHIJKLMNOPQRST":
            ws[f"{col}{r}"].border = Border(bottom=Side(style="thin", color=NAVY))
widths(ws, {c: 13 for c in "ABCDEFGHIJKLMNOPQRST"})
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 8
ws.freeze_panes = "B6"
printsetup(ws, title_rows="4:5")
FPA = lambda col: q(FP, f"${col}${R0}:${col}${R1}")

# =====================================================================
# ESCENARIOS
# =====================================================================
ws = S["Escenarios"]
ES = "Escenarios"
title(ws, "Escenarios de comisión y tiempo de recuperación",
      "Casa completa con 7 habitaciones arrendadas y lavandería con monedero. La comisión de administración se calcula sobre el ingreso neto operativo (después de gastos).")
section(ws, "A4", "1. Inversión total a recuperar")
inv = [
    (5, "Obra y adecuaciones (incluye pendientes, bomba y 4x1000)", f'=SUMIFS({M_I},{M_M},"Obra y adecuaciones",{M_L},"Gasto")', "link"),
    (6, "Mobiliario 7 habitaciones (real + estimado)", f"={sp('$I$13')}", "link"),
    (7, "Lavandería (lavadora + monedero)", f"={sp('$B$45')}", "link"),
    (8, "Subtotal capital invertido (menos venta de objetos almacenados)", f"=SUM(B5:B7)-{sp('$F$28')}", "f"),
    (9, "Intereses préstamo mamá (36 cuotas)", f"={q(PR, '$B$13')}", "link"),
    (10, "INVERSIÓN TOTAL A RECUPERAR", "=B8+B9", "f"),
]
for r, lab, f_, kind in inv:
    put(ws, f"A{r}", lab, bold=r in (8, 10))
    put(ws, f"B{r}", f_, kind, fmt=CUR, bold=r in (8, 10))
section(ws, "A14", "2. Casa completa: 7 habitaciones + lavandería (estabilizado)")
header(ws, 15, ["Concepto", "Comisión 10%", "Comisión 20%", "Comisión 30%"])
for c, prow in (("B", "$B$29"), ("C", "$B$30"), ("D", "$B$31")):
    ws[f"{c}15"].value = f"={sp(prow)}"
    ws[f"{c}15"].number_format = '"Comisión "0%'
pc = {"B": sp("$B$29"), "C": sp("$B$30"), "D": sp("$B$31")}
cum = {"B": "R", "C": "S", "D": "T"}
rows = [
    (16, "Ingreso bruto mensual (arriendos + lavandería)", lambda c: f"={sp('$C$13')}+{sp('$F$30')}*{sp('$F$29')}*{sp('$F$32')}*COUNTA({sp('$A$6:$A$12')})", CUR),
    (17, "Gastos de operación (extrapolados a 7 inquilinos, con ahorro de la lavadora)", lambda c: f"=-({sp('$B$22')}+({sp('$C$22')}-{sp('$B$22')})*(COUNTA({sp('$A$6:$A$12')})-1)/({sp('$F$33')}-1)+{sp('$C$23')}*COUNTA({sp('$A$6:$A$12')})/{sp('$F$33')}+{sp('$C$24')}-{sp('$B$47')})", CUR),
    (18, "Neto operativo", lambda c: f"={c}16+{c}17", CUR),
    (19, "Comisión de administración", lambda c: f"=-MAX(0,{c}18)*{pc[c]}", CUR),
    (20, "Neto casa para los abuelos (antes de deudas)", lambda c: f"={c}18+{c}19", CUR),
    (21, "Cuota préstamo mamá (dic 2026 – nov 2029)", lambda c: f"=-{q(PR, '$B$11')}", CUR),
    (22, "Neto para los abuelos mientras se paga el préstamo", lambda c: f"={c}20+{c}21", CUR),
    (23, "Neto anual (antes de deudas)", lambda c: f"={c}20*12", CUR),
    (24, "Rentabilidad anual sobre la inversión total", lambda c: f"=IFERROR({c}23/$B$10,0)", PCT),
    (25, "Recuperación simple, como si la casa estuviera llena desde el día 1 (años)", lambda c: f"=IFERROR($B$10/{c}23,0)", '0.0'),
    (26, "Recuperación con llenado gradual (meses desde oct 2026)",
     lambda c: f'=IF(COUNTIF({FPA(cum[c])},"<"&$B$10)>=72,"más de 72",COUNTIF({FPA(cum[c])},"<"&$B$10)+1)', '0'),
    (27, "Mes en que se recupera toda la inversión", lambda c: f'=IF(ISNUMBER({c}26),EDATE({sp("$B$54")},{c}26-1),"—")', MON),
    (28, "Ganancia adicional vs. antes de la remodelación (mensual)", lambda c: f"={c}20-{sp('$B$50')}", CUR),
    (29, "Ingreso mensual de Felipe por administración", lambda c: f"=-{c}19", CUR),
]
for r, lab, fn, fmt in rows:
    put(ws, f"A{r}", lab, bold=r in (20, 22, 24, 26))
    for c in "BCD":
        put(ws, f"{c}{r}", fn(c), fmt=fmt, bold=r in (20, 22, 24, 26), align="right")
section(ws, "A32", "3. Referencia de mercado para la comisión")
put(ws, "A33", "Honorario inmobiliaria tradicional: % sobre canon bruto + IVA")
put(ws, "B33", f"={sp('$B$52')}*(1+{sp('$B$53')})", "link", fmt=PCT)
put(ws, "A34", "Equivale en pesos a (mensual, casa completa)")
put(ws, "B34", "=B33*B16", fmt=CUR)
put(ws, "A35", "Equivale como % del ingreso neto operativo")
put(ws, "B35", "=IFERROR(B34/B18,0)", fmt=PCT, bold=True)
put(ws, "A36", "Una inmobiliaria administra 1 contrato por inmueble. Aquí son 7 contratos, servicios compartidos, rotación de inquilinos, aseo y mantenimiento de un coliving: más carga que una administración tradicional, y la inmobiliaria además suele cobrar un canon por cada inquilino nuevo.", "note", wrap=True)
ws.merge_cells("A36:D38")
put(ws, "A40", "Fuente referencia: Inmobiliare Latam, 'Administración de inmuebles en Bogotá: guía 2026' (8%–12% del canon + IVA).", "note", italic=True)
section(ws, "A42", "4. Mientras la habitación 7 sigue de bodega (6 habitaciones + lavandería)")
header(ws, 43, ["Concepto", "Comisión 10%", "Comisión 20%", "Comisión 30%"])
for c, prow in (("B", "$B$29"), ("C", "$B$30"), ("D", "$B$31")):
    ws[f"{c}43"].value = f"={sp(prow)}"
    ws[f"{c}43"].number_format = '"Comisión "0%'
rows6 = [
    (44, "Ingreso bruto mensual (arriendos + lavandería)", lambda c: f"=SUM({sp('$C$6:$C$11')})+{sp('$F$30')}*{sp('$F$29')}*{sp('$F$32')}*6"),
    (45, "Gastos de operación (con ahorro de la lavadora)", lambda c: f"=-({sp('$C$25')}-{sp('$B$47')})"),
    (46, "Neto operativo", lambda c: f"={c}44+{c}45"),
    (47, "Comisión de administración", lambda c: f"=-MAX(0,{c}46)*{pc[c]}"),
    (48, "Neto casa para los abuelos (antes de deudas)", lambda c: f"={c}46+{c}47"),
    (49, "Cuota préstamo mamá", lambda c: f"=-{q(PR, '$B$11')}"),
    (50, "Neto para los abuelos mientras se paga el préstamo", lambda c: f"={c}48+{c}49"),
]
for r, lab, fn in rows6:
    put(ws, f"A{r}", lab, bold=r in (48, 50))
    for c in "BCD":
        put(ws, f"{c}{r}", fn(c), fmt=CUR, bold=r in (48, 50), align="right")
widths(ws, {"A": 64, "B": 18, "C": 18, "D": 18})
printsetup(ws, landscape=False)

# =====================================================================
# COMPARACIÓN
# =====================================================================
ws = S["Comparación"]
CP = "Comparación"
title(ws, "Comparación: la casa frente a otras alternativas",
      "Rentabilidades antes de impuestos. Los CDT y la cuenta de ahorros tienen retención en la fuente; la casa ya descuenta predial y reservas.")
section(ws, "A4", "1. ¿Cuánto rinde el mismo dinero en cada alternativa? (capital = inversión total)")
header(ws, 5, ["Alternativa", "Rentabilidad anual", "Ingreso anual", "Ingreso mensual promedio"])
alts = [
    (6, "Cuenta de ahorros actual de los abuelos", f"={sp('$B$38')}"),
    (7, "CDT 360 días — promedio del mercado (sep 2026)", f"={sp('$B$40')}"),
    (8, "CDT 360 días — mejor tasa del mercado (sep 2026)", f"={sp('$B$41')}"),
    (9, "La Clarita — comisión 10%", f"={q(ES, '$B$24')}"),
    (10, "La Clarita — comisión 20%", f"={q(ES, '$C$24')}"),
    (11, "La Clarita — comisión 30%", f"={q(ES, '$D$24')}"),
]
for r, lab, f_ in alts:
    put(ws, f"A{r}", lab, bold=r >= 9)
    put(ws, f"B{r}", f_, "link", fmt=PCT)
    put(ws, f"C{r}", f"=B{r}*{q(ES, '$B$10')}", fmt=CUR)
    put(ws, f"D{r}", f"=C{r}/12", fmt=CUR)
put(ws, "A12", "La casa rinde más, pero exige administración, tiene riesgo de vacancia y el capital queda en el inmueble (no se puede retirar como un CDT). A cambio, la casa conserva su valor como propiedad.", "note", wrap=True)
ws.merge_cells("A12:D13")

section(ws, "A15", "2. Opción A vs. Opción B — ¿de qué caja salen los muebles?")
header(ws, 16, ["Concepto", "A 6 meses", "A 12 meses", "A 24 meses"])
put(ws, "A17", "Fecha de corte")
for c, m in (("B", 5), ("C", 11), ("D", 23)):
    put(ws, f"{c}17", f"=EDATE({sp('$B$54')},{m})", fmt=MON, align="right")
cmp2 = [(18, "Caja de la casa acumulada — Opción A (arriendo paga muebles)", "M"),
        (19, "Caja de la casa acumulada — Opción B (ahorros pagan muebles)", "O"),
        (20, "Dinero tomado de los ahorros para muebles y lavadora — Opción B", "Q")]
for r, lab, col in cmp2:
    put(ws, f"A{r}", lab, bold=r == 19)
    for c in "BCD":
        put(ws, f"{c}{r}", f"=INDEX({FPA(col)},MATCH({c}$17,{FPA('A')},0))", "link", fmt=CUR, bold=r == 19)
put(ws, "A21", "Saldo ILUSTRATIVO de ahorros después de pagar los muebles (sin contar intereses)")
for c in "BCD":
    put(ws, f"{c}21", f"={sp('$B$39')}-{c}20", fmt=CUR)
put(ws, "A22", "En las dos opciones los abuelos terminan con prácticamente el mismo patrimonio: los muebles se pagan con la misma plata. La diferencia es que en la Opción A la casa no alcanza a cubrir sus obligaciones mes a mes y alguien tiene que tapar el hueco.", "note", wrap=True)
ws.merge_cells("A22:D23")

section(ws, "A25", "3. Señales de la Opción A (el arriendo paga préstamo + muebles)")
put(ws, "A26", "Meses con caja negativa en los primeros 12 meses")
put(ws, "B26", f'=COUNTIF({q(FP, f"$L${R0}:$L${R0 + 11}")},"<0")', "link", fmt='0')
put(ws, "A27", "Peor momento de la caja acumulada (primeros 24 meses)")
put(ws, "B27", f"=MIN({q(FP, f'$M${R0}:$M${R0 + 23}')})", "link", fmt=CUR)
put(ws, "A28", "Resultado de caja del mes más difícil")
put(ws, "B28", f"=MIN({q(FP, f'$L${R0}:$L${R0 + 23}')})", "link", fmt=CUR)
put(ws, "A29", "Ese faltante no tiene fuente: si no sale de los ahorros, se deja de pagar la cuota a mamá o la tarjeta, y no queda ganancia para nadie.", "note", wrap=True)
ws.merge_cells("A29:D30")

section(ws, "A32", "4. Si el efectivo se guarda para otra propiedad")
put(ws, "A33", "Rendimiento extra al año por cada $10.000.000 que pasen de la cuenta (8,5%) a un CDT promedio")
put(ws, "B33", f"=10000000*({sp('$B$40')}-{sp('$B$38')})", fmt=CUR)
put(ws, "A34", "Fuentes: Infobae (29 sep 2026) y El País (sep 2026) — promedio CDT 360 días 12,07% EA según Banco de la República; finanzasplus.co — tasas hasta 13% EA.", "note", italic=True, wrap=True)
ws.merge_cells("A34:D35")
widths(ws, {"A": 66, "B": 18, "C": 18, "D": 18})
printsetup(ws, landscape=False)

# =====================================================================
# VALORACIÓN
# =====================================================================
ws = S["Valoración"]
VA = "Valoración"
title(ws, "¿Cuánto podría valer la casa para un inversionista?",
      "Estimación de referencia por dos métodos. Antes de publicar un precio: avalúo comercial profesional (Lonja o perito RAA).")
section(ws, "A4", "Datos del inmueble")
dat = [(5, "Barrio / localidad", "La Clarita, Engativá — Bogotá", None),
       (6, "Estrato", 3, '0'),
       (7, "Frente del lote (m)", 6, '0.0'),
       (8, "Fondo del lote (m)", 12, '0.0'),
       (9, "Área del lote (m²)", "=B7*B8", '0'),
       (10, "Área construida (m²)", 190, '0'),
       (11, "Pisos", "3 + terraza", None),
       (12, "Habitaciones de renta", f"=COUNTA({sp('$A$6:$A$12')})", '0')]
for r, lab, v, fmt in dat:
    put(ws, f"A{r}", lab)
    put(ws, f"B{r}", v, "in" if not str(v).startswith("=") else "f", fmt=fmt)

section(ws, "A14", "Método 1 — Comparables en venta en La Clarita (precios de oferta)")
header(ws, 15, ["Inmueble", "Área construida (m²)", "Precio de oferta", "Precio por m²"])
comps = [("Casa 4 hab / 3 baños", 155, 450000000), ("Casa 4 hab / 2 baños", 120, 550000000),
         ("Casa 4 hab / 2 baños", 242, 580000000), ("Casa 3 hab / 3 baños", 160, 640000000)]
for i, (n, a, p) in enumerate(comps):
    r = 16 + i
    put(ws, f"A{r}", n)
    put(ws, f"B{r}", a, "in", fmt='0')
    put(ws, f"C{r}", p, "in", fmt=CUR)
    put(ws, f"D{r}", f"=C{r}/B{r}", fmt=CUR)
put(ws, "A20", "Promedio comparables", bold=True)
put(ws, "D20", "=AVERAGE(D16:D19)", bold=True, fmt=CUR)
put(ws, "A21", "Promedio Engativá (casas, precio por m²)")
put(ws, "D21", 3412157, "in", fmt=CUR)
put(ws, "A22", "Descuento de negociación sobre precio de oferta")
put(ws, "D22", 0.07, "in", fmt=PCT)
put(ws, "A23", "Valor por comparables (promedio La Clarita, con descuento)", bold=True)
put(ws, "D23", "=B10*D20*(1-D22)", bold=True, fmt=CUR)
put(ws, "A24", "Valor por comparables (promedio Engativá, con descuento)", bold=True)
put(ws, "D24", "=B10*D21*(1-D22)", bold=True, fmt=CUR)
put(ws, "A25", "Fuentes: listados en Fincaraíz / Properati / Mercado Libre para La Clarita, y Habi.co ('Valor del metro cuadrado en Bogotá por localidad'), consultados 30 sep 2026. Son precios de oferta, no de cierre. Los comparables no están remodelados para renta.", "note", italic=True, wrap=True)
ws.merge_cells("A25:D26")

section(ws, "A28", "Método 2 — Por la renta que produce (lo que mira un inversionista)")
header(ws, 28, ["Método 2 — Por la renta que produce (lo que mira un inversionista)", "Casa completa (7 hab)", "Hoy (6 hab)"], height=32)
ws["A28"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
m2 = [(29, "Ingreso neto operativo anual estabilizado (arriendos + lavandería − gastos)", f"=({q(ES, '$B$18')})*12", f"=({q(ES, '$B$46')})*12"),
      (30, "Menos administración a precio de mercado (el comprador tendría que pagarla)", f"=-{q(ES, '$B$34')}*12", f"=-{q(ES, '$B$33')}*{q(ES, '$B$44')}*12"),
      (31, "Ingreso neto anual para un inversionista", "=B29+B30", "=C29+C30"),
      (32, "Equivale a un flujo de caja mensual de", "=B31/12", "=C31/12")]
for r, lab, f7, f6 in m2:
    put(ws, f"A{r}", lab, bold=r in (31, 32))
    put(ws, f"B{r}", f7, "link" if "!" in f7 else "f", fmt=CUR, bold=r in (31, 32))
    put(ws, f"C{r}", f6, "link" if "!" in f6 else "f", fmt=CUR, bold=r in (31, 32))
header(ws, 33, ["Tasa de capitalización exigida", "Valor — casa completa", "Valor — hoy (6 hab)", "Lectura"])
caps = [(34, 0.07, "Inversionista que valora la renta ya probada"), (35, 0.08, "Punto medio"),
        (36, 0.09, "Inversionista que castiga el riesgo del coliving")]
for r, cr, n in caps:
    put(ws, f"A{r}", cr, "in", fmt=PCT, align="center")
    put(ws, f"B{r}", f"=B$31/A{r}", fmt=CUR, bold=True)
    put(ws, f"C{r}", f"=C$31/A{r}", fmt=CUR)
    put(ws, f"D{r}", n, "note")
put(ws, "A37", "Tasa de capitalización = ingreso neto anual ÷ precio. Supuesto de Felipe/Claude, no dato de mercado: para vivienda por habitaciones con 7 contratos se exige más que para un arriendo tradicional. La venta se entrega amoblada y con contratos vigentes.", "note", italic=True, wrap=True)
ws.merge_cells("A37:D38")

section(ws, "A40", "Rango de referencia")
put(ws, "A41", "Valor mínimo de los métodos", bold=True)
put(ws, "B41", "=MIN(D23,D24,B34,B35,B36)", bold=True, fmt=CUR)
put(ws, "A42", "Valor máximo de los métodos", bold=True)
put(ws, "B42", "=MAX(D23,D24,B34,B35,B36)", bold=True, fmt=CUR)
put(ws, "A43", "Punto medio", bold=True)
put(ws, "B43", "=AVERAGE(B41,B42)", bold=True, fmt=CUR)
put(ws, "A45", "Qué revisaría un comprador: (0) un comprador paga por ingreso con historial, no por ingreso proyectado — la Hab 7 y el monedero aún no existen, y el Registro Mensual es la evidencia que respalda el valor por renta; (1) el ingreso depende de que alguien administre las 7 habitaciones; (2) la casa es una sola unidad, sin independización; (3) la obra fue interna sin licencia — conviene confirmar con un curador o abogado que las adecuaciones (baños y cocinas nuevas) no la requerían; (4) contratos vigentes y su historial de pago.", "note", wrap=True)
ws.merge_cells("A45:D49")
widths(ws, {"A": 60, "B": 20, "C": 20, "D": 40})
printsetup(ws, landscape=False)

# =====================================================================
# CUADRE
# =====================================================================
ws = S["Cuadre"]
CU = "Cuadre"
title(ws, "Fuentes y usos del dinero — cuadre",
      "Todo sale de la hoja Movimientos. Si agrega un pago allá, este cuadre se actualiza solo.")
section(ws, "A4", "1. ¿En qué se gastó y quién lo pagó? (incluye pagos pendientes)")
src_cols = "BCDEFG"
header(ws, 5, ["Categoría"] + fuentes + ["Total"], height=40)
for i, (c_, g) in enumerate(cats[:-1]):
    r = 6 + i
    put(ws, f"A{r}", f"={sp(f'$A${58 + i}')}", "link")
    for j, col in enumerate(src_cols):
        put(ws, f"{col}{r}", f'=SUMIFS({M_I},{M_C},$A{r},{M_F},{col}$5,{M_L},"Gasto")', fmt=CUR)
    put(ws, f"H{r}", f"=SUM(B{r}:G{r})", fmt=CUR, bold=True)
tr = 6 + len(cats) - 1  # 13
put(ws, f"A{tr}", "TOTAL", bold=True)
for col in src_cols + "H":
    put(ws, f"{col}{tr}", f"=SUM({col}6:{col}{tr - 1})", fmt=CUR, bold=True)
    ws[f"{col}{tr}"].fill = PatternFill("solid", start_color=LIGHT)
for j, col in enumerate(src_cols):
    ws[f"{col}5"].value = f"={sp(f'$D${58 + j}')}"
put(ws, f"A{tr + 1}", "Control: total cuadre = total Movimientos")
put(ws, f"H{tr + 1}", f'=IF(ROUND(H{tr}-{q(MV, "$I$4")},0)=0,"OK","REVISAR")', bold=True)

section(ws, "A17", "2. Préstamo de mamá — ¿a dónde fue cada peso?")
put(ws, "A18", "Monto recibido")
put(ws, "B18", f"={sp('$B$32')}", "link", fmt=CUR)
put(ws, "A19", "Pagado a proveedores (contratista y cámaras)")
put(ws, "B19", f'=SUMIFS({M_I},{M_F},"Préstamo mamá",{M_L},"Gasto",{M_C},"<>Costos bancarios – 4x1000")', fmt=CUR)
put(ws, "C19", "Aparte: 4x1000 asumido por la cuenta de mamá, adicional al préstamo:", "note")
put(ws, "H19", f'=SUMIFS({M_I},{M_F},"Préstamo mamá",{M_C},"Costos bancarios – 4x1000")', fmt=CUR)
put(ws, "A20", "Trasladado a la cuenta de la abuela")
put(ws, "B20", f'=SUMIFS({M_I},{M_F},"Préstamo mamá",{M_L},"Traslado")', fmt=CUR)
put(ws, "A21", "Diferencia sin explicar (debe ser 0)", bold=True)
put(ws, "B21", "=B18-B19-B20", bold=True, fmt=CUR)

section(ws, "A23", "3. Contratista principal — contratado vs. pagado")
put(ws, "A24", "Contrato inicial")
put(ws, "B24", 85000000, "in", fmt=CUR)
put(ws, "A25", "Acuerdo de arreglos adicionales")
put(ws, "B25", 29000000, "in", fmt=CUR)
put(ws, "A26", "Total contratado", bold=True)
put(ws, "B26", "=B24+B25", bold=True, fmt=CUR)
put(ws, "A27", "Total pagado (12 pagos: 5 recibos en efectivo + 7 transferencias)")
put(ws, "B27", f'=SUMIFS({M_I},{M_C},"Obra – contratista principal",{M_L},"Gasto")', fmt=CUR)
put(ws, "A28", "Saldo por pagar al contratista", bold=True)
put(ws, "B28", "=B26-B27", bold=True, fmt=CUR)
put(ws, "C24", "Abuelo: $84.500.000 · Préstamo mamá: $29.500.000", "note")

section(ws, "A30", "4. Pagado vs. pendiente")
put(ws, "A31", "Ya pagado")
put(ws, "B31", f'=SUMIFS({M_I},{M_L},"Gasto",{M_K},"<>Pendiente de pago")', fmt=CUR)
put(ws, "A32", "Pendiente de pago (saldo eléctrico + mano de obra bomba)")
put(ws, "B32", f'=SUMIFS({M_I},{M_L},"Gasto",{M_K},"Pendiente de pago")', fmt=CUR)
put(ws, "A33", "Total", bold=True)
put(ws, "B33", "=B31+B32", bold=True, fmt=CUR)

section(ws, "A35", "5. Tarjeta de crédito de los abuelos")
put(ws, "A36", "Compras cargadas a la tarjeta (mobiliario + bomba)")
put(ws, "B36", f'=SUMIFS({M_I},{M_F},"TC abuelos (cuotas)",{M_L},"Gasto")', fmt=CUR)
put(ws, "A37", "Cuotas pagadas (según Registro Mensual)")
put(ws, "B37", f"=SUM({q('Registro Mensual', '$K$6:$K$41')})", "link", fmt=CUR)
put(ws, "A38", "Saldo pendiente en la tarjeta", bold=True)
put(ws, "B38", "=B36-B37", bold=True, fmt=CUR)

section(ws, "A40", "6. Pendiente: aportes de Felipe")
put(ws, "A41", "Pagos registrados desde la cuenta de Felipe con fuente 'Felipe (por definir)'")
put(ws, "B41", f'=SUMIFS({M_I},{M_F},"Felipe (por definir)",{M_L},"Gasto")', fmt=CUR)
put(ws, "A42", "Además, Felipe pagó con su cuenta y tarjetas otras compras de la casa. Se documentarán con extractos bancarios y se definirá si son reembolso o aporte.", "note", wrap=True)
ws.merge_cells("A42:H43")
widths(ws, {"A": 58, "B": 16, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 16})
printsetup(ws)

# =====================================================================
# REGISTRO MENSUAL (live)
# =====================================================================
ws = S["Registro Mensual"]
RM = "Registro Mensual"
title(ws, "Registro mensual real — documento vivo",
      "Cada mes: escriba en azul lo que realmente entró y salió. La columna 'Proyectado' compara contra el plan.")
header(ws, 5, ["Mes", "Hab. ocupadas", "Arriendos cobrados", "Servicios pagados", "Reserva apartada",
               "Predial apartado", "Neto operativo", "Comisión administración", "Neto casa (antes de deudas)",
               "Cuota préstamo pagada", "Cuotas tarjeta pagadas con arriendo", "Neto entregado a los abuelos",
               "Neto proyectado (antes de deudas)", "Diferencia real − proyectado", "Notas"], height=44)
for k in range(36):
    r = 6 + k
    put(ws, f"A{r}", dt.date(2026, 9, 1) if k == 0 else f"=EDATE(A{r - 1},1)", "in" if k == 0 else "f", fmt=MON)
    for col in "BCDEFJK":
        style(ws[f"{col}{r}"], "in", fmt=('0' if col == "B" else CUR))
    put(ws, f"G{r}", f'=IF(C{r}="","",C{r}-D{r}-E{r}-F{r})', fmt=CUR)
    put(ws, f"H{r}", f'=IF(G{r}="","",MAX(0,G{r})*{sp("$B$28")})', fmt=CUR)
    put(ws, f"I{r}", f'=IF(G{r}="","",G{r}-H{r})', fmt=CUR, bold=True)
    put(ws, f"L{r}", f'=IF(I{r}="","",I{r}-J{r}-K{r})', fmt=CUR, bold=True)
    put(ws, f"M{r}", f"=IFERROR(INDEX({FPA('I')},MATCH(A{r},{FPA('A')},0)),\"\")", "link", fmt=CUR)
    put(ws, f"N{r}", f'=IF(OR(I{r}="",M{r}=""),"",I{r}-M{r})', fmt=CUR)
    style(ws[f"O{r}"], "in", wrap=True)
    for col in "ABCDEFGHIJKLMNO":
        ws[f"{col}{r}"].border = BOX
# September 2026 — real example row
ws["B6"].value = 1
ws["C6"].value = 1250000
ws["D6"].value = 606000
ws["E6"].value = 0
ws["F6"].value = 0
ws["J6"].value = 0
ws["K6"].value = 1190248
ws["O6"].value = "Arriendo Hab 1 se usó completo para la 1ª cuota de la tarjeta (muebles). Servicios = valores actuales reportados. Comisión causada, no cobrada. Nadie recibió ganancia."
confirm(ws["D6"], "Servicios de sep: se usaron los valores 'actuales' que reportó Felipe. Reemplazar con los recibos reales.")
widths(ws, {"A": 10, "B": 8, "C": 13, "D": 13, "E": 12, "F": 12, "G": 13, "H": 13, "I": 14, "J": 13, "K": 14, "L": 14, "M": 14, "N": 14, "O": 44})
ws.freeze_panes = "B6"
printsetup(ws, title_rows="5:5")

# =====================================================================
# PROPUESTA
# =====================================================================
ws = S["Propuesta"]
title(ws, "Propuesta: separar la inversión del flujo mensual de la casa",
      "Preparado por Felipe Fajardo para Jorge Quintero · corte 30 de septiembre de 2026")
r = 4


def para(text, rows=2, bold=False):
    global r
    c = put(ws, f"A{r}", text, wrap=True, bold=bold)
    ws.merge_cells(f"A{r}:D{r + rows - 1}")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += rows


def kv(label, formula, fmt=CUR, kind="f", bold=False):
    global r
    put(ws, f"A{r}", label, bold=bold)
    ws.merge_cells(f"A{r}:C{r}")
    put(ws, f"D{r}", formula, kind, fmt=fmt, bold=bold)
    r += 1


selE = lambda row: f"=INDEX({q(ES, f'$B${row}:$D${row}')},MATCH({sp('$B$28')},{q(ES, '$B$15:$D$15')},0))"

section(ws, f"A{r}", "1. Lo que pasó en septiembre")
r += 1
kv("Arriendo cobrado Habitación 1", f"={q(RM, '$C$6')}")
kv("Primera cuota de la tarjeta (muebles de la Habitación 1)", f"=-{q(RM, '$K$6')}")
kv("Servicios del mes", f"=-{q(RM, '$D$6')}")
kv("Resultado del mes", f"={q(RM, '$C$6')}-{q(RM, '$K$6')}-{q(RM, '$D$6')}", bold=True)
para("Después de pagar la primera cuota de los muebles y los servicios, el mes cerró en negativo. Desde el segundo mes cada habitación amoblada suma una cuota de ~$850.000. Si cada habitación nueva se amuebla a cuotas y se paga con el arriendo, y desde diciembre además se paga el préstamo de mamá con el mismo arriendo, la casa no alcanza a cubrir sus obligaciones durante el primer año y no deja ganancia para nadie.", 3)
r += 1
section(ws, f"A{r}", "2. El principio: dos cajas separadas")
r += 1
para("Caja de INVERSIÓN (se paga una sola vez): remodelación, muebles, lavadora. La obra se pagó con ahorros y con el préstamo, no con el arriendo. Los muebles son parte de esa misma inversión.", 2)
para("Caja de OPERACIÓN (cada mes): arriendos menos servicios, reservas y administración. De aquí sale la cuota del préstamo de mamá, y lo que sobra es la ganancia real de la casa.", 2)
r += 1
section(ws, f"A{r}", "3. Si el arriendo paga también los muebles (Opción A)")
r += 1
kv("Cuota mensual del préstamo de mamá desde el 15 dic 2026", f"={q(PR, '$B$11')}")
kv("Meses con caja negativa en los primeros 12 meses", f"={q(CP, '$B$26')}", fmt='0')
kv("Peor momento acumulado de la caja de la casa", f"={q(CP, '$B$27')}", bold=True)
para("Ese faltante tendría que salir de algún lado. Si no sale de los ahorros, se atrasa la cuota a mamá o la de la tarjeta. Es decir, la plata de los muebles termina saliendo de los ahorros de todas formas, pero tarde y en desorden.", 3)
r += 1
section(ws, f"A{r}", "4. Lo que propongo (Opción B)")
r += 1
para("a) Pagar las cuotas de los muebles y de la lavandería desde la cuenta de ahorros, no desde el arriendo. Se siguen aprovechando las cuotas sin interés de Mercado Libre: no hay que pagar de contado, solo cambiar de qué cuenta sale cada cuota.", 3)
para("b) El arriendo paga la operación de la casa y la cuota del préstamo de mamá. Desde que se arriendan las 6 habitaciones alcanza para todo y sobra; con la 7 sobra más.", 2)
kv("Ganancia mensual mientras se paga el préstamo — 6 habitaciones", selE(50), bold=True)
kv("Ganancia mensual mientras se paga el préstamo — casa completa (7)", selE(22))
kv("Ganancia mensual cuando termine el préstamo — casa completa (7)", selE(20))
para("c) Cambiar la lavadora de más de 20 años por una nueva e instalar un monedero para cobrar por ciclo. El monedero ordena el uso de la máquina entre 7 inquilinos y convierte la lavandería en un ingreso. El contrato de arriendo ya contempla el cobro por ciclo; hoy el uso es gratis hasta que se instale. El monedero se compra a cuotas sin interés.", 3)
kv("Lavadora nueva", f"={q('Mobiliario', '$C$52')}")
kv("Monedero", f"={q('Mobiliario', '$B$54')}*{q('Mobiliario', '$B$56')}")
kv("Ingreso mensual estimado del monedero (casa completa)", f"={sp('$F$30')}*{sp('$F$29')}*{sp('$F$32')}*COUNTA({sp('$A$6:$A$12')})")
kv("Ahorro mensual en energía y agua (lavadora nueva)", f"={sp('$B$47')}")
kv("Meses para que la lavandería completa se pague sola", f"=IFERROR({sp('$B$45')}/(D{r-2}+D{r-1}),0)", fmt='0', bold=True)
para("d) El dinero que se quiera guardar para otra propiedad rinde más en un CDT a 1 año (~12% EA promedio del mercado) que en la cuenta actual (8,5%).", 2)
kv("Rendimiento extra al año por cada $10.000.000", f"={q(CP, '$B$33')}")
para("e) Habilitar la habitación 7 cuando salgan los objetos almacenados (se trasladan a la nueva propiedad o se venden). Si se venden, lo recibido ayuda a pagar el mobiliario de esa habitación. Los objetos probablemente valen menos de $10.000.000, y cada mes de bodega deja de producir la ganancia de abajo: en un año, mantener la bodega cuesta aproximadamente lo mismo que valen los objetos.", 3)
kv("Canon esperado habitación 7", f"={sp('$C$12')}")
kv("Fecha estimada", f"={sp('$D$12')}", fmt=MON)
kv("Mobiliario habitación 7 (estimado)", f"={q('Mobiliario', '$H$44')}")
kv("Ganancia mensual adicional para los abuelos con la habitación 7", selE(20) + "-" + selE(48)[1:], bold=True)
kv("Lo que deja de producir la bodega en 12 meses", "=12*(" + selE(20)[1:] + "-" + selE(48)[1:] + ")")
r += 1
section(ws, f"A{r}", "5. Qué cambia y qué no")
r += 1
kv("Dinero neto que saldría de los ahorros para muebles y lavandería", f"={q(FP, '$Q$77')}")
kv("Caja de la casa a 12 meses — Opción A", f"={q(CP, '$C$18')}")
kv("Caja de la casa a 12 meses — Opción B", f"={q(CP, '$C$19')}", bold=True)
para("La rentabilidad total de la inversión es la misma en las dos opciones. Lo que cambia es que con la Opción B la casa cumple sus obligaciones cada mes, la cuota a mamá está asegurada desde el primer pago y la ganancia real de la casa se ve desde el primer mes.", 3)
r += 1
section(ws, f"A{r}", "6. Decisión pendiente: comisión de administración")
r += 1
para("Felipe administra las 7 habitaciones: contratos, cobros, servicios, mantenimiento, rotación de inquilinos y marketing. Se presentan tres opciones sobre el ingreso neto (después de gastos). La decisión es de los abuelos. Ver hoja Escenarios.", 3)
ws.column_dimensions["A"].width = 48
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 20
printsetup(ws, landscape=False)

# =====================================================================
# RESUMEN
# =====================================================================
ws = S["Resumen"]
title(ws, "Casa La Clarita — Rendición de cuentas e inversión",
      "Preparado por Felipe Fajardo para Jorge Quintero · corte 30 de septiembre de 2026 · cifras en pesos colombianos")
r = 4


def blk(t):
    global r
    section(ws, f"A{r}", t)
    ws[f"A{r}"].border = Border(bottom=Side(style="thin", color=NAVY))
    ws[f"B{r}"].border = Border(bottom=Side(style="thin", color=NAVY))
    r += 1


def line(label, formula, fmt=CUR, bold=False, kind="f"):
    global r
    put(ws, f"A{r}", label, bold=bold)
    put(ws, f"B{r}", formula, kind, fmt=fmt, bold=bold, align="right")
    if bold:
        ws[f"A{r}"].fill = ws[f"B{r}"].fill = PatternFill("solid", start_color=LIGHT)
    r += 1


cat = lambda name: f'=SUMIFS({M_I},{M_C},"{name}",{M_L},"Gasto")'
blk("1. ¿Cuánto se ha invertido?")
line("Contratista principal (contrato $85M + adicionales $29M)", cat("Obra – contratista principal"))
line("Adecuación eléctrica y aumento de carga (incluye saldo pendiente)", cat("Eléctrico – aumento de carga (Enel)"))
line("Cámaras de seguridad", cat("Seguridad – cámaras"))
line("Bomba de presión (equipo + instalación)", cat("Hidráulico – bomba de presión"))
line("Impuesto 4x1000 de transferencias", cat("Costos bancarios – 4x1000"))
line("Mobiliario habitación 1", cat("Mobiliario – habitaciones"))
line("TOTAL INVERTIDO Y COMPROMETIDO", f"={q(MV, '$I$4')}", bold=True)
line("     Ya pagado", f"={q(CU, '$B$31')}")
line("     Pendiente de pago", f"={q(CU, '$B$32')}")
r += 1
blk("2. ¿De dónde salió el dinero?")
for f_ in ["Abuelos – capital propio", "Préstamo mamá", "TC abuelos (cuotas)"]:
    line(f_, f'=SUMIFS({M_I},{M_F},"{f_}",{M_L},"Gasto")')
line("Sobrante del préstamo en la cuenta de la abuela", f"={q(CU, '$B$20')}")
r += 1
blk("3. Lo que falta por invertir (proyectado)")
line("Mobiliario habitaciones 2 a 7 (estimado)", f"={sp('$I$13')}-{q('Mobiliario', '$B$44')}")
line("Lavandería (lavadora + monedero)", f"={sp('$B$45')}")
line("Menos: venta de objetos almacenados", f"=-{sp('$F$28')}")
line("Intereses del préstamo de mamá (36 cuotas)", f"={q(PR, '$B$13')}")
line("INVERSIÓN TOTAL A RECUPERAR", f"={q(ES, '$B$10')}", bold=True)
r += 1
blk("4. ¿Qué va a producir la casa? (casa completa: 7 habitaciones + lavandería)")
sel = lambda row: f"=INDEX({q(ES, f'$B${row}:$D${row}')},MATCH({sp('$B$28')},{q(ES, '$B$15:$D$15')},0))"
line("Ingreso bruto mensual (arriendos + lavandería)", f"={q(ES, '$B$16')}")
line("Gastos de operación", f"={q(ES, '$B$17')}")
line("Comisión de administración", sel(19))
line("Comisión seleccionada (sobre ingreso neto)", f"={sp('$B$28')}", fmt=PCT)
line("GANANCIA MENSUAL PARA LOS ABUELOS", sel(20), bold=True)
line("Cuota préstamo de mamá (dic 2026 – nov 2029)", f"={q(ES, '$B$21')}")
line("Ganancia mensual mientras se paga el préstamo", sel(22))
line("Mientras la hab. 7 sea bodega (6 hab.), mientras se paga el préstamo", sel(50))
line("Antes de la remodelación", f"={sp('$B$50')}")
r += 1
blk("5. ¿En cuánto tiempo se recupera?")
line("Meses para recuperar toda la inversión (desde oct 2026)", sel(26), fmt='0', bold=True)
line("Mes estimado de recuperación total", sel(27), fmt=MON)
line("Rentabilidad anual de la casa", sel(24), fmt=PCT, bold=True)
line("vs. CDT a 1 año (promedio del mercado)", f"={sp('$B$40')}", fmt=PCT)
line("vs. cuenta de ahorros actual", f"={sp('$B$38')}", fmt=PCT)
r += 1
blk("6. ¿Cuánto podría valer la casa? (referencia, no avalúo)")
line("Rango bajo", f"={q(VA, '$B$41')}")
line("Rango alto", f"={q(VA, '$B$42')}")
r += 1
put(ws, f"A{r}", "Detalle de cada pago en la hoja Movimientos. Supuestos editables en la hoja Supuestos. Celdas amarillas = por confirmar.", "note", italic=True)
widths(ws, {"A": 62, "B": 22})
printsetup(ws, landscape=False)

wb.save(OUT)
print("saved", OUT)
