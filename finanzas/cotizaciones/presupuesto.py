import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Crear libro de trabajo
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Presupuesto Auditado"

# Mostrar líneas de cuadrícula
ws.views.sheetView[0].showGridLines = True

# Paleta de Colores
NAVY_HEADER = "1B365D"
WHITE = "FFFFFF"
ZEBRA_FILL = "F9FAFC"
BORDER_GRAY = "D9D9D9"
SUBTOTAL_CASO_BG = "D5E1F2"
SUBTOTAL_EXTRA_BG = "FFE6B3"

# Estilos de Fuentes
font_title = Font(name="Calibri", size=16, bold=True, color="1B365D")
font_subtitle = Font(name="Calibri", size=10, italic=True, color="555555")
font_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_bold = Font(name="Calibri", size=10, bold=True)
font_regular = Font(name="Calibri", size=10)
font_subtotal = Font(name="Calibri", size=11, bold=True, color="000000")
font_total = Font(name="Calibri", size=12, bold=True, color=WHITE)

align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center")
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(left=Side(style='thin', color=BORDER_GRAY), right=Side(style='thin', color=BORDER_GRAY), top=Side(style='thin', color=BORDER_GRAY), bottom=Side(style='thin', color=BORDER_GRAY))
double_bottom_border = Border(left=Side(style='thin', color=BORDER_GRAY), right=Side(style='thin', color=BORDER_GRAY), top=Side(style='thin', color=BORDER_GRAY), bottom=Side(style='double', color="FFFFFF"))

CURRENCY_FORMAT = '"$"#,##0'

# Encabezado
ws.merge_cells("A1:E1")
ws["A1"] = "PRESUPUESTO DE REMODELACIÓN Y EXTRAS (VERSIÓN AUDITADA Y AJUSTADA)"
ws["A1"].font = font_title

ws.merge_cells("A2:E2")
ws["A2"] = "Desglose comparativo ajustado tras auditoría de insumos, cantidades de cableado, impermeabilización y morteros (COP)"
ws["A2"].font = font_subtitle

headers = ["Concepto (Nombre del Trabajo)", "Rango Inferior\n(Barato / Económico)", "Rango Medio\n(Justo / Mercado Real)", "Rango Superior\n(Precio Inflado)", "Detalle del Trabajo Auditado\n(Materiales y Proceso de Mano de Obra)"]

ws.row_dimensions[4].height = 32
for col_num, header_title in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_num, value=header_title)
    cell.font = font_header
    cell.fill = PatternFill(start_color=NAVY_HEADER, fill_type="solid")
    cell.alignment = align_header
    cell.border = thin_border

# Datos auditados y corregidos
data = [
    ("Caso 1: 5 Baños Nuevos + Remodelaciones", 35100000, 44870000, 60500000, "MATERIALES: Red hidrosanitaria PVC agua fría/caliente/reventilación, enchapes cerámicos (18m² por baño), 7 divisiones en vidrio templado 6mm, sanitarios, lavamanos con mueble, griferías, impermeabilizante de zonas húmedas (Sika/Bronco) y 5 rejillas/sifones antiolor.\nMANO DE OBRA: Regatado, tuberías, impermeabilización de duchas, enchapado, fraguado, montaje de sanitarios, griferías, divisiones y carpintería."),
    ("Caso 2: 5 Cocinas Integrales (1.20m)", 9900000, 14950000, 23500000, "MATERIALES: 5 muebles integrales (aéreo y bajo) de 1.20m en RH, pocetas inox, grifería, mangueras/conexiones de gas flexible, estufas empotrables y campanas extractoras.\nMANO DE OBRA: Armado y nivelación de muebles, adecuación de acometidas hidráulicas y de gas, empotrado de estufas, conexión y prueba de hermeticidad de campanas."),
    ("Caso 3: Acabados Generales", 5950000, 9100000, 14350000, "MATERIALES: Pintura tipo 1 lavable, estuco plástico, lijas, placas de Drywall 8mm/12mm, perfiles metálicos, resanes de fachada y punto de gas.\nMANO DE OBRA: Preparación y lijado de muros, aplicación de 3 manos de pintura, armado de estructura y muros en Drywall, resane exterior y mantenimiento en terraza."),
    ("Extra 1: Reparación Piso Terraza (2 m²)", 240000, 440000, 760000, "MATERIALES: Manto/solución impermeabilizante de alto tráfico, mortero de nivelación, cerámica exterior y boquilla hidrorepelente.\nMANO DE OBRA: Demolición de baldosa fracturada, retirado de escombro, aplicación de capa impermeabilizante, vaciado de mortero e instalación de cerámica."),
    ("Extra 2: Puertas, Mueble y 7 Chapas", 885000, 1355000, 2250000, "MATERIALES: 7 chapas/cerraduras de manija moderna tipo pomo/palanca, masilla para madera, lijas y herrajes.\nMANO DE OBRA: Cepillado y ajuste de hojas de madera desalineadas, reparación estructural de mueble de madera, desmonte de chapas antiguas e instalación de las 7 nuevas."),
    ("Extra 3: Estructura Metálica + Tanque 500L", 1080000, 1750000, 2850000, "MATERIALES: Perfiles de acero anticorrosivo, tanque de reserva de 500L, tubería PVC presión, válvulas de corte y flotador.\nMANO DE OBRA: Corte, soldadura y armado de base metálica, fijación e instalación del tanque, trazado de tubería de alimentación y empalme a red hidráulica."),
    ("Extra 4: Mantenimiento Marquesina (9 m²)", 200000, 350000, 550000, "MATERIALES: Desengrasante industrial, limpiador de juntas, silicona neutra de alta adherencia e intemperie.\nMANO DE OBRA: Lavado profundo de estructura y vidrios/policarbonato, retiro de silicona cristalizada/viejas juntas y resellado perimetral e interuniones."),
    ("Extra 5: Cambio de 4 Vidrios Rotos", 180000, 300000, 520000, "MATERIALES: 4 láminas de vidrio transparente (4mm a 5mm) cortadas a medida, silicona de sellado y empaques.\nMANO DE OBRA: Retiro con cuidado de junquillos y esquirlas de vidrio roto, limpieza de marcos, asentado de cristales nuevos y sellado perimetral."),
    ("Extra 6: 2 Chapas Yale Puerta Principal", 440000, 610000, 850000, "MATERIALES: 2 cerraduras de seguridad de sobreponer tipo Yale (1 con sistema de alarma por fuerza/vibración) y tornillería pesada.\nMANO DE OBRA: Perforación, avellanado y calado en puerta principal de madera/metal, montaje de cilindros, escudos de seguridad y prueba de mecanismos."),
    ("Extra 7: Impermeabilización Manto a Calor", 380000, 530000, 800000, "MATERIALES: Emulsión asfáltica de imprimación, manto asfáltico de 3mm/4mm y gas propano para soplete.\nMANO DE OBRA: Imprimación de la superficie en canaleta/muro, calentamiento a llama de soplete y traslape de manto asfáltico sellado al vacío."),
    ("Extra 8: Cielo Raso PVC (24 m² - 2 Hab)", 1210000, 1740000, 2650000, "MATERIALES: Tablilla PVC machihembrada, perfilería galvanizada (omega/ángulo), cornisas, chazos de expansión/puntillas de impacto y conectores eléctricos.\nMANO DE OBRA: Trazado de nivel, armado de estructura colgada, ensamble de tablillas PVC, remate perimetral con cornisas y reconexión de puntos de luz."),
    ("Extra 9: Reubicación Eléctrica (24 Puntos)", 1880000, 2750000, 4200000, "MATERIALES: 45m tubería PVC conduit 1/2\", cajas 4x2, 140m de cable de cobre THHN #12 AWG (Fase/Neutro/Tierra), tomas dobles con placa, mortero y estuco.\nMANO DE OBRA: Regatado en muro, tendido de tubería, cableado completo de circuitos, conexión de tomas y resane de regatas."),
    ("Extra 10: 15 Luminarias LED Sobreponer", 525000, 750000, 1275000, "MATERIALES: 15 plafones LED de sobreponer (12W-24W) con drivers integrados, chazos y conectores eléctricos.\nMANO DE OBRA: Desmonte de rosetas antiguas, chazado de platinas a placa/cielo raso, empalme eléctrico y montaje estético de lámparas."),
    ("Extra 11: Independización por Sondeo", 600000, 950000, 1500000, "MATERIALES: Cinta aislante vulcanizada, guía sonda de fontanero/electricista, lubricante de cables y marcadores.\nMANO DE OBRA: Mapeo de circuitos, retiro de líneas cruzadas entre P1 y P2, sondeo mediante guía en tuberías empotradas y recableado e independización en tablero."),
    ("Extra 12: Marco + Puerta Habitación P1", 320000, 480000, 750000, "MATERIALES: Marco en madera ajustado, hoja entamborada triplex con moldura, 3 bisagras 3.5\", chapa de pomo y tapaluces.\nMANO DE OBRA: Presentación e instalación de marco a plomo, fijación de bisagras, colgado de hoja de puerta, ajuste de chapa y tapaluces."),
    ("Extra 13: Red 220V Bifásica (7 Duchas)", 3490000, 4310000, 6050000, "MATERIALES: 450m cable THHN #10 AWG, 150m tubería conduit 3/4\", 7 breakers bifásicos (2x30A/40A) y caja subdistribuidora.\nMANO DE OBRA: Tendido de 7 circuitos independientes a 220V desde tablero principal, peinado de caja de brekers y conexión directa a duchas."),
    ("Extra 14: Cambio de Piso Cerámica (18 m²)", 1380000, 2050000, 3100000, "MATERIALES: Cerámica de alto tráfico, pegante Pegacor (4-5 bultos), boquilla, guardaescobas cerámicos y mortero de nivelación.\nMANO DE OBRA: Demolición de baldosa previa, cargue y retiro de escombros, alistado con mortero, asentado de cerámica, fraguado e instalación de guardaescobas.")
]

current_row = 5
for row_idx, item in enumerate(data):
    concepto, econ, justo, caro, detalle = item
    ws.row_dimensions[current_row].height = 55 if len(detalle) > 150 else 38
    bg_fill = PatternFill(start_color=(ZEBRA_FILL if row_idx % 2 == 1 else WHITE), fill_type="solid")
    
    ws.cell(row=current_row, column=1, value=concepto).font = font_bold
    ws.cell(row=current_row, column=5, value=detalle).font = font_regular
    
    for c_idx, val in zip([2, 3, 4], [econ, justo, caro]):
        cell = ws.cell(row=current_row, column=c_idx, value=val)
        cell.number_format = CURRENCY_FORMAT
        cell.font = font_bold if c_idx == 3 else font_regular
        cell.alignment = align_right

    for c in range(1, 6):
        ws.cell(row=current_row, column=c).fill = bg_fill
        ws.cell(row=current_row, column=c).border = thin_border
    current_row += 1

# Subtotal Base
ws.row_dimensions[current_row].height = 24
ws.cell(row=current_row, column=1, value="SUBTOTAL CASOS BASE (Auditado)").font = font_subtotal
for c in range(1, 6):
    cell = ws.cell(row=current_row, column=c)
    cell.fill = PatternFill(start_color=SUBTOTAL_CASO_BG, fill_type="solid")
    cell.border = thin_border
    if 2 <= c <= 4:
        col_let = get_column_letter(c)
        cell.value = f"=SUM({col_let}5:{col_let}7)"
        cell.number_format = CURRENCY_FORMAT
        cell.font = font_subtotal
        cell.alignment = align_right
current_row += 1

# Subtotal Extras
ws.row_dimensions[current_row].height = 24
ws.cell(row=current_row, column=1, value="SUBTOTAL TRABAJOS EXTRAS (Auditado)").font = font_subtotal
for c in range(1, 6):
    cell = ws.cell(row=current_row, column=c)
    cell.fill = PatternFill(start_color=SUBTOTAL_EXTRA_BG, fill_type="solid")
    cell.border = thin_border
    if 2 <= c <= 4:
        col_let = get_column_letter(c)
        cell.value = f"=SUM({col_let}8:{col_let}21)"
        cell.number_format = CURRENCY_FORMAT
        cell.font = font_subtotal
        cell.alignment = align_right
current_row += 1

# Total General
tot_row = current_row
ws.row_dimensions[tot_row].height = 30
ws.cell(row=tot_row, column=1, value="TOTAL GENERAL ACUMULADO (AUDITADO)").font = font_total
for c in range(1, 6):
    cell = ws.cell(row=tot_row, column=c)
    cell.fill = PatternFill(start_color=NAVY_HEADER, fill_type="solid")
    cell.border = double_bottom_border
    if 2 <= c <= 4:
        col_let = get_column_letter(c)
        cell.value = f"=SUM({col_let}5:{col_let}21)"
        cell.number_format = CURRENCY_FORMAT
        cell.font = font_total
        cell.alignment = align_right

ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 24
ws.column_dimensions["D"].width = 24
ws.column_dimensions["E"].width = 85

wb.save("Presupuesto_Remodelacion_Auditado.xlsx")